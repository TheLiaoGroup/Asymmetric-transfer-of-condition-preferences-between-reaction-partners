"""Figure 1 and its Source Data, generated together from the public measurements.

Every plotted value is recomputed here from the released measurements with the analysis code's own estimators
(06_CODE/Core/00_common/ctb/anatomy.py) and must agree with the approved analysis outputs to 1e-9 before anything is
drawn. Two values come from the approved outputs because they need the full analysis run: the 95% interval of the
median molecular share (2,000 hierarchical resamples) and the owner named by the likelihood estimator; rebuild_all.py step 13
recomputes both from the measurements (06_CODE/Checks/frozen_reads_recomputed.py).

Public measurements, 05_DATA/Release/00_RAW_IMMUTABLE/DATASET_analysis.csv.gz:
  A1  numeric yields of the 15 arrays, Hiyama C46 excluded (177,308 rows): panels b-d
  A3  Hiyama C46, the duplicate of C43 (74 rows): noise calibration of panel c
Rows of the process-development (CMC) screens are dropped as the file is read; no other file of measurements is opened.
Approved outputs, 05_DATA/Results/display_inputs: Figure1b_example_selection.json (fixed example rule),
Figure1b_example_residuals.csv, Anatomy_arrays.csv, Anatomy.json. Structures: 02_FIGURES/ChemDraw (chemdraw_structures.py).

Writes 02_FIGURES/Main/Figure1.{pdf,svg,png}, 04_SOURCE_DATA/Figure1/Figure1_SourceData.xlsx and
Figure1_SourceData_provenance.json, and 08_REVIEW_AND_QC/figure1/figure1_plot_readback.json (what each panel draws,
read back from the drawn artists).
    python 06_CODE/Figures/figure1.py             all outputs in place
    python 06_CODE/Figures/figure1.py --out DIR   all outputs into DIR, for comparison with the delivered files
"""
import argparse
import csv
import hashlib
import itertools
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                                              # noqa: E402
from matplotlib.colors import Normalize, to_hex                              # noqa: E402
from matplotlib.lines import Line2D                                          # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle   # noqa: E402
from matplotlib.text import Text                                             # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / "06_CODE/Core/00_common/ctb"))
import fig_style as fs                                                       # noqa: E402
import structure_art as sa                                                   # noqa: E402
from anatomy import FORMS, Shares, electrophile_role, noise_variance, pattern     # noqa: E402
from history import ROLE                                                     # noqa: E402

RAW = ROOT / "05_DATA/Release/00_RAW_IMMUTABLE/DATASET_analysis.csv.gz"
DI = ROOT / "05_DATA/Results/display_inputs"
APPROVED = {n: DI / n for n in ("Figure1b_example_selection.json", "Figure1b_example_residuals.csv",
                                "Anatomy_arrays.csv", "Anatomy.json")}
CHEMDRAW = ROOT / "02_FIGURES/ChemDraw"
SPEC = ROOT / "07_METHODS_AND_PROVENANCE/FIGURE_STRUCTURES_PROMOTION_READY.csv"
OUT = dict(pdf=ROOT / "02_FIGURES/Main/Figure1.pdf", svg=ROOT / "02_FIGURES/Main/Figure1.svg",
           png=ROOT / "02_FIGURES/Main/Figure1.png", xlsx=ROOT / "04_SOURCE_DATA/Figure1/Figure1_SourceData.xlsx",
           provenance=ROOT / "04_SOURCE_DATA/Figure1/Figure1_SourceData_provenance.json",
           readback=ROOT / "08_REVIEW_AND_QC/figure1/figure1_plot_readback.json")
N_A1, N_A3 = 177308, 74
TOL = 1e-9
CLIP = 30.0             # colour range of panel b, percentage points
LABEL_PT = 7.0          # atom-label size of the panel-a structures; one scale for all four
PNG_DPI = 1200
NAME = {  # names used in the Figure 1 legend and the panel lock
    "CNOC": "N,O-dimethylhydroxylamine", "O=C(O)c1ccco1": "furan-2-carboxylic acid",
    "O=C(O)c1ccc(Br)s1": "5-bromothiophene-2-carboxylic acid", "CNc1ccc(F)cc1": "4-fluoro-N-methylaniline",
    "Cc1ccc(F)cc1C(=O)O": "5-fluoro-2-methylbenzoic acid", "N#Cc1ccc(C(=O)O)cc1": "4-cyanobenzoic acid"}
TITLE = dict(a="Replacing one partner tests what the other carries",
             b="Four couplings of one amine share a condition pattern",
             c="About a third of the pattern belongs to the molecules",
             d="One partner carries most of the molecular share")
TRANSFORM = dict(
    a="Schematic. A-D are molecules of the amide coupling array (SMILES in this sheet); A + B, A + C and D + B are "
      "product-forming pairs of that array (A1). No measurement is plotted.",
    b="A1 rows of the amide coupling array (Acylation|1). A pair is product-forming if its highest yield over the "
      "conditions is at least 0.10. Condition mean = mean yield of a condition over the product-forming pairs. Each "
      "pair's yields are fitted by least squares to intercept + slope x condition mean over its 94 conditions; condition "
      "pattern = 100 x (yield - fitted yield), in percentage points. Cell colour = pattern clipped to +/-30 points. "
      "r = Pearson correlation of two rows' patterns over the 94 conditions; mean r = mean of the six r within a group.",
    c="For each of the 15 arrays (A1 rows), condition patterns as in panel b over the array's product-forming pairs. "
      "Share of role 1 (2) = mean, over ordered pairs of pairs that share only their role-1 (role-2) molecule, of the "
      "mean product of the two patterns over the conditions measured for both, divided by the mean over pairs of the "
      "mean squared pattern. Molecular = role 1 + role 2. Noise = mean over pairs of the mean model measurement variance "
      "left in a pattern, k x max(0.07 x clip(fitted, 0, 1), 0.012)^2 x (1 - leverage), divided by the same mean squared "
      "pattern; k makes the model reproduce the Hiyama C43/C46 duplicate (C43 from A1, C46 from A3). Unattributed "
      "= 1 - molecular - noise (pair-specific interaction and any error outside the noise model). Neither shared = the role-share ratio for pairs sharing neither molecule. All x 100 (%).",
    d="Role shares of panel c (x 100, %). Owner = role with the larger share; owner fraction = owner share / (role 1 + "
      "role 2 share); owner >= 2x partner compares the owner share with twice the partner share, a negative partner share "
      "counted as zero. Electrophilic halide = the role named '... halide' (10 arrays). Likelihood owner: approved output "
      "of the Gaussian likelihood fit (anatomy.ml_components), not rerun here.")
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
rel = lambda p: Path(p).resolve().relative_to(ROOT).as_posix()
neg = lambda s: s.replace("-", "−")                     # typographic minus in printed numbers


# ====================================================================== public measurements
def public_units():
    cols = ["screen", "condition_id", "subpair", "sub_1_smiles", "sub_2_smiles", "yield"]
    A = pd.read_csv(RAW, usecols=cols, low_memory=False)
    A = A[~A.screen.str.contains("CMC", regex=False)]                         # process-development screens: not used
    y = pd.to_numeric(A["yield"], errors="coerce")
    c46 = (A.screen == "Hiyama|1") & (A.condition_id == "C46")
    A1 = A[y.notna() & ~c46].assign(**{"yield": y[y.notna() & ~c46]})
    A3 = A[c46].assign(**{"yield": y[c46]})
    assert len(A1) == N_A1 and len(A3) == N_A3 and A3["yield"].notna().all(), (len(A1), len(A3))
    assert set(A1.screen) == set(fs.ARRAY_LABEL)
    assert (A1.subpair == A1.sub_1_smiles + " >> " + A1.sub_2_smiles).all()
    return A1, A3


def matrix(g):
    P = g.pivot(index="subpair", columns="condition_id", values="yield")
    a = np.array([p.split(" >> ")[0] for p in P.index]); b = np.array([p.split(" >> ")[1] for p in P.index])
    return P, a, b


def noise_scale(A1, A3):
    """anatomy.hiyama_noise_scale on the released rows: Hiyama C43 from A1, its duplicate C46 from A3."""
    P, _, _ = matrix(A1[A1.screen == "Hiyama|1"])
    d46 = A3.set_index("subpair")["yield"].reindex(P.index).to_numpy(float)
    M = P.to_numpy(float); E, H, fit = pattern(M); j = list(P.columns).index("C43")
    ok = np.isfinite(d46) & np.isfinite(M[:, j])
    e43, e46 = E[ok, j], d46[ok] - fit[ok, j]
    r = np.mean(e43 * e46) / np.mean((e43 ** 2 + e46 ** 2) / 2)
    n_model = np.mean(np.maximum(0.07 * np.clip(fit[ok, j], 0, 1), 0.012) ** 2 * (1 - H[ok, j]))
    return float(r), float((1 - r) * np.mean((e43 ** 2 + e46 ** 2) / 2) / n_model), int(ok.sum()), float(np.corrcoef(e43, e46)[0, 1])


# ====================================================================== panel values
def panel_a(A1, sel):
    spec = [r for r in csv.DictReader(open(SPEC, encoding="utf-8")) if r["figure"] == "1"]
    S = {r["label"].split()[0]: r["smiles"] for r in spec}
    assert S["A"] == sel["shared_molecule"] and [S["B"], S["C"]] == sel["shared_partners"][:2]
    man = {r["name"]: r for r in csv.DictReader(open(CHEMDRAW / "structures_manifest.csv", encoding="utf-8"))}
    art = {}
    for L in "ABCD":
        m = man[f"Fig1_a_{L}"]
        assert m["smiles"] == S[L], L
        art[L] = sa.load(CHEMDRAW / f"Fig1_a_{L}.svg")
        assert art[L]["sha256"] == m["svg_sha256"], f"{L}: SVG differs from the ChemDraw manifest"
    amide = A1[A1.screen == sel["screen"]]
    for l, r in (("A", "B"), ("A", "C"), ("D", "B")):                         # the three pairs drawn are product-forming pairs
        y = amide[amide.subpair == f"{S[l]} >> {S[r]}"]["yield"]
        assert len(y) and y.max() >= FORMS, (l, r)
    role = {L: ROLE[sel["screen"]][0][0] if L in "AD" else ROLE[sel["screen"]][1][0] for L in "ABCD"}
    shown = dict(A="A + B (central pair); A + C (A kept)", B="A + B (central pair); D + B (B kept)",
                 C="A + C (replaces B)", D="D + B (replaces A)")
    rows = [dict(letter=L, name=NAME[S[L]], smiles=S[L], role=role[L], pairs_shown=shown[L],
                 chemdraw_file=f"Fig1_a_{L}.cdxml") for L in "ABCD"]
    return pd.DataFrame(rows), art, S


def panel_b(A1, sel, frozen, letter_of):
    key = sel["screen"]
    P, a, b = matrix(A1[A1.screen == key])
    M = P.to_numpy(float); act = np.nanmax(M, 1) >= FORMS
    E, H, fit = pattern(M[act])
    s = np.nanmean(M[act], 0)
    pf = list(P.index[act])
    assert sel["role"] == "1"                                                  # the shared molecule is partner 1
    shared = [f"{sel['shared_molecule']} >> {x}" for x in sel["shared_partners"]]
    order = shared + list(sel["unrelated_pairs"])
    conds = sorted(P.columns, key=lambda c: int(c[1:]))
    jc = [list(P.columns).index(c) for c in conds]
    rows = [pf.index(p) for p in order]
    Z = 100 * E[rows][:, jc]
    assert frozen.pair.tolist() == order and list(frozen.columns[2:]) == conds and np.isfinite(Z).all()
    dz = float(np.abs(Z - frozen[conds].to_numpy(float)).max())
    assert dz < TOL, dz
    Y = M[act][rows][:, jc]; S = s[jc]
    coef = np.array([np.linalg.lstsq(np.c_[np.ones(len(S)), S], y, rcond=None)[0] for y in Y])
    F = coef[:, :1] + coef[:, 1:] * S
    assert np.abs(F - fit[rows][:, jc]).max() < 1e-12 and np.abs(100 * (Y - F) - Z).max() < TOL
    # comparison pairs: which molecules they share, read from the pairs themselves
    n1 = [p.split(" >> ")[0] for p in order]; n2 = [p.split(" >> ")[1] for p in order]
    share_cmp = [(i, j) for i, j in itertools.combinations(range(4, 8), 2) if n1[i] == n1[j] or n2[i] == n2[j]]
    assert share_cmp == []                                                      # the comparison pairs share no molecule (2026-10-13)
    assert not ({n1[i] for i in range(4, 8)} | {n2[i] for i in range(4, 8)}) & ({n1[0]} | set(n2[:4]))
    r = {}
    for g, ix in (("shared", range(4)), ("comparison", range(4, 8))):
        r[g] = {(i, j): float(np.corrcoef(Z[i], Z[j])[0, 1]) for i, j in itertools.combinations(ix, 2)}
    mean_r = {g: float(np.mean(list(v.values()))) for g, v in r.items()}
    assert abs(mean_r["shared"] - sel["mean_r_shared"]) < 1e-12 and abs(mean_r["comparison"] - sel["mean_r_unrelated"]) < 1e-12

    def label(i):
        if i < 4:
            acid = sel["shared_partners"][i]
            return r"$\mathbf{A}$ + " + (rf"$\mathbf{{{letter_of[acid]}}}$" if acid in letter_of else NAME[acid])
        return f"comparison {i - 3}"
    labels = [label(i) for i in range(8)]
    plain = [l.replace(r"$\mathbf{", "").replace("}$", "") for l in labels]
    cells = [dict(group="shared A" if i < 4 else "comparison", row=i + 1, row_label=plain[i], n_h_partner=n1[i],
                  carboxylic_acid=n2[i], condition_id=c, column=k + 1, yield_fraction=Y[i, k], condition_mean_fraction=S[k],
                  pair_intercept=coef[i, 0], pair_slope=coef[i, 1], fitted_yield_fraction=F[i, k], pattern_points=Z[i, k],
                  colour_value_points=float(np.clip(Z[i, k], -CLIP, CLIP)))
             for i in range(8) for k, c in enumerate(conds)]
    corr = [dict(group=g, row_i=i + 1, row_j=j + 1, label_i=plain[i], label_j=plain[j], r=v)
            for g in r for (i, j), v in r[g].items()]
    corr += [dict(group=g, row_i=None, row_j=None, label_i="mean of the six r", label_j=None, r=mean_r[g]) for g in r]
    corr = pd.DataFrame(corr).astype({"row_i": "Int64", "row_j": "Int64"})
    info = dict(screen=key, pairs=int(len(M)), product_forming=int(act.sum()), conditions=len(conds),
                a1_rows_array=int((A1.screen == key).sum()), a1_rows_plotted=int(np.isfinite(Y).sum()), max_abs_diff_points=dz)
    return dict(Z=Z, labels=labels, conds=conds, r=r, mean_r=mean_r, cells=pd.DataFrame(cells), corr=corr, info=info)


def panels_cd(A1, A3, approved_arrays, J):
    r_dup, k, n_dup, pearson_dup = noise_scale(A1, A3)
    assert abs(r_dup - J["share"]["hiyama_duplicate_reliability"]) < TOL and abs(k - J["share"]["noise_scale"]) < TOL
    rows = []
    for key in sorted(fs.ARRAY_LABEL):
        g = A1[A1.screen == key]; P, a, b = matrix(g); M = P.to_numpy(float)
        act = np.nanmax(M, 1) >= FORMS
        E, H, fit = pattern(M[act]); sh = Shares(E, a[act], b[act]); s = sh.compute(np.ones(act.sum()))
        noise = noise_variance(M[act], fit, H, k) / sh.var.mean()
        er = electrophile_role(key)
        own = "1" if s["1"] >= s["2"] else "2"; par = "2" if own == "1" else "1"
        rows.append(dict(screen=key, array=fs.ARRAY_LABEL[key], a1_rows=len(g), conditions=M.shape[1], pairs=len(M),
                         product_forming=int(act.sum()), role1=ROLE[key][0][0], role2=ROLE[key][1][0],
                         share1=s["1"], share2=s["2"], neither=s["N"], molecular=s["1"] + s["2"], noise=noise,
                         combination=1 - s["1"] - s["2"] - noise, owner_role=own, owner=ROLE[key][int(own) - 1][0],
                         partner=ROLE[key][int(par) - 1][0], owner_share=s[own], partner_share=s[par],
                         owner_fraction=s[own] / (s["1"] + s["2"]), twice=bool(s[own] >= 2 * max(s[par], 0.0)),
                         halide=None if er is None else ROLE[key][int(er) - 1][0],
                         halide_owns=None if er is None else bool(s[er] > s["2" if er == "1" else "1"])))
    T = pd.DataFrame(rows).set_index("screen")
    ref = approved_arrays.set_index("screen").loc[T.index]
    diffs = {}
    for mine, theirs in (("pairs", "pairs"), ("product_forming", "product_forming"), ("share1", "share1"),
                         ("share2", "share2"), ("neither", "share_neither"), ("molecular", "share"), ("noise", "noise"),
                         ("combination", "combination"), ("owner_share", "owner_share"),
                         ("partner_share", "partner_share"), ("owner_fraction", "owner_fraction")):
        diffs[mine] = float(np.abs(T[mine].to_numpy(float) - ref[theirs].to_numpy(float)).max())
        assert diffs[mine] < TOL, (mine, diffs[mine])
    assert (T.owner_role.astype(int) == ref.owner_role).all() and (T.owner == ref.owner).all()
    hal = T.halide.notna()
    assert (T.halide_owns[hal].astype(bool) == ref.electrophile_carries[hal].astype(bool)).all()
    T["ml_owner_role"] = ref.ml_owner_role.astype(int).astype(str)          # approved output, not rerun
    T["ml_same_owner"] = T.ml_owner_role == T.owner_role
    med = {c: float(T[c].median()) for c in ("molecular", "combination", "noise", "neither")}
    sj = J["share"]
    for c, jk in (("molecular", "median"), ("combination", "combination_median"), ("noise", "noise_median"),
                  ("neither", "neither_median")):
        assert abs(med[c] - sj[jk]) < TOL, c
    own = dict(median_owner_fraction=float(T.owner_fraction.median()), twice=int(T.twice.sum()),
               same_owner=int(T.ml_same_owner.sum()), halide_arrays=int(hal.sum()),
               halide_owns=int(T.halide_owns[hal].astype(bool).sum()), arrays=len(T))
    o = J["ownership"]
    assert (abs(own["median_owner_fraction"] - o["median_owner_fraction"]) < TOL and own["twice"] == o["owner_at_least_twice_partner"]
            and own["same_owner"] == o["same_owner_both_estimators"] and own["arrays"] == o["arrays"] == J["n_arrays"]
            and own["halide_arrays"] == J["electrophile"]["arrays"] and own["halide_owns"] == J["electrophile"]["carries"])
    assert int(T.pairs.sum()) == J["pairs"] and int(T.product_forming.sum()) == J["product_forming_pairs"]
    halide_first, others = fs.array_order({key: electrophile_role(key) is not None for key in T.index})
    T = T.loc[halide_first + others]
    T["group"] = ["electrophilic halide" if h else "other" for h in T.halide.notna()]
    T["plot_row"] = np.arange(1, len(T) + 1)
    info = dict(noise_calibration=dict(duplicate_pairs=n_dup, duplicate_agreement=r_dup, duplicate_pattern_pearson=pearson_dup, k=k), max_abs_diff=diffs,
                arrays=len(T), pairs=int(T.pairs.sum()), product_forming=int(T.product_forming.sum()),
                a1_rows=int(T.a1_rows.sum()))
    return T, med, own, list(sj["ci"]), info


# ====================================================================== drawing
W = fs.WIDTH_MM
H = 157.0
YB, YC = 29.0, 76.0                                    # tops of the panel-b and panel-c/d rows
HX0, HX1, ROWH = 41.5, 149.5, 3.0                      # panel-b heatmap columns and row height
BLOCK = dict(shared=YB + 9.1, comparison=YB + 25.5)    # tops of the two heatmap blocks
YKEY = YB + 39.4                                       # panel-b line under the heatmap: axis note and colour key
RX0, RW, R0 = 152.5, 17.0, -0.4                        # panel-b r strip: left edge, mm per unit r, r at the left edge
CX0, CW, DX0, DW = 27.5, 54.0, 104.0, 48.0             # panel c and d axes (mm)
ROW0, PITCH, GROUP_GAP = YC + 10.0, 4.0, 3.0           # first array row top, row pitch, gap between the groups


def row_y(i):
    return ROW0 + PITCH / 2 + PITCH * i + (GROUP_GAP if i >= 10 else 0.0)


def draw(sel, a_rows, art, B, T, med, own, ci):
    fs.apply(); fs.check_font()
    plt.rcParams["svg.hashsalt"] = "figure1"
    fig = plt.figure(figsize=(W * fs.MM, H * fs.MM))
    cv = fig.add_axes([0, 0, 1, 1]); cv.set_xlim(0, W); cv.set_ylim(H, 0); cv.axis("off"); cv.set_gid("canvas")
    put = lambda x, y, s, gid=None, **kw: cv.text(x, y, s, gid=gid, **kw)
    for L, x, y in (("a", 0, 0), ("b", 0, YB), ("c", 0, YC), ("d", DX0 - 8.0, YC)):
        fs.letter(fig, x, y, L, H).set_gid(f"letter|{L}")
        fs.title(fig, x + 4.2, y + 0.35, TITLE[L], H).set_gid(f"title|{L}")

    # ---------------------------------------------------------------- a: reciprocal replacement
    sz = {L: sa.size_mm(art[L], LABEL_PT) for L in "ABCD"}
    gap = 2.4
    width = lambda l, r: sz[l][0] + 2 * gap + sz[r][0]
    yc = 12.5
    ylet = yc + max(h for _, h in sz.values()) / 2 + 2.0
    xc = W / 2 - width("A", "B") / 2
    arm = 22.0
    pos = dict(left=(xc - arm - width("A", "C"), ("A", "C"), {"A"}), centre=(xc, ("A", "B"), {"A", "B"}),
               right=(xc + width("A", "B") + arm, ("D", "B"), {"B"}))
    tile = dict(A=fs.C["keepA"], B=fs.C["keepB"])
    for tag, (x0, (l, r), keep) in pos.items():
        for L, x in ((l, x0), (r, x0 + sz[l][0] + 2 * gap)):
            w, h = sz[L]; y0 = yc - h / 2
            if L in keep:
                cv.add_patch(FancyBboxPatch((x - 1.0, y0 - 1.0), w + 2.0, h + 2.0, boxstyle="round,pad=0,rounding_size=1.3",
                                            facecolor=tile[L], edgecolor="none", zorder=1, gid=f"a|{tag}|tile|{L}"))
            sa.draw(cv, art[L], x, y0, LABEL_PT, gid=f"a|{tag}|{L}")
            put(x + w / 2, ylet, L, f"a|{tag}|letter|{L}", fontsize=7, fontweight="bold", ha="center", va="top")
        put(x0 + sz[l][0] + gap, yc, "+", f"a|{tag}|plus", fontsize=7, ha="center", va="center")
    for tag, x_from, x_to, kept, swapped in (("left", xc - 2.2, pos["left"][0] + width("A", "C") + 2.2, "A", "B"),
                                             ("right", xc + width("A", "B") + 2.2, pos["right"][0] - 2.2, "B", "A")):
        cv.add_patch(FancyArrowPatch((x_from, yc), (x_to, yc), arrowstyle="-|>", mutation_scale=6, lw=0.6,
                                     color=fs.C["muted"], shrinkA=0, shrinkB=0, gid=f"a|arrow|{tag}"))
        put((x_from + x_to) / 2, yc - 1.4, rf"keep $\mathbf{{{kept}}}$, replace $\mathbf{{{swapped}}}$", f"a|arrow_label|{tag}",
            fontsize=fs.PT["small"], ha="center", va="bottom")
        x0, (l, r), _ = pos[tag]
        put(x0 + width(l, r) / 2, ylet + 3.6, rf"measures what $\mathbf{{{kept}}}$ carries", f"a|outcome|{tag}",
            fontsize=fs.PT["small"], ha="center", va="top")

    # ---------------------------------------------------------------- b: measured condition patterns
    norm = Normalize(-CLIP, CLIP)
    cmap = plt.get_cmap("RdBu_r")
    header = dict(shared=r"Four couplings sharing $\mathbf{A}$",
                  comparison="Four comparison pairs sharing no molecule")
    for g, ix in (("shared", slice(0, 4)), ("comparison", slice(4, 8))):
        top = BLOCK[g]
        put(HX0, top - 2.3, header[g], f"b|header|{g}", fontsize=fs.PT["small"], color=fs.C["muted"], ha="left", va="center")
        ax = fs.axes_mm(fig, HX0, top, HX1 - HX0, 4 * ROWH, H)
        im = ax.imshow(B["Z"][ix], cmap=cmap, norm=norm, aspect="auto", interpolation="nearest", extent=(0, 94, 4, 0))
        im.set_gid(f"b|heatmap|{g}")
        for yy in (1, 2, 3):
            ax.axhline(yy, color="white", lw=0.6)
        ax.set_xlim(0, 94); ax.set_ylim(4, 0); ax.set_xticks([])
        ax.set_yticks([0.5, 1.5, 2.5, 3.5], B["labels"][ix], fontsize=fs.PT["small"])
        ax.tick_params(axis="y", length=0, pad=1.5)
        for sp in ax.spines.values():
            sp.set_linewidth(fs.LW["thin"]); sp.set_color(fs.C["muted"])
        ax.set_gid(f"b|axes|{g}")
        ym = top + 2 * ROWH
        put(RX0, ym - 1.5, rf"mean $r$ = {B['mean_r'][g]:.2f}", f"b|mean_r|{g}", fontsize=fs.PT["base"], ha="left", va="center")
        ys = ym + 2.3
        rx = lambda v: RX0 + (v - R0) * RW
        assert min(B["r"][g].values()) >= R0
        cv.plot([rx(R0), rx(1)], [ys, ys], color=fs.C["faint"], lw=0.4, solid_capstyle="butt", gid=f"b|r_axis|{g}")
        for v, s in ((0, "0"), (1, "1")):
            cv.plot([rx(v), rx(v)], [ys - 0.5, ys + 0.5], color=fs.C["faint"], lw=0.4, gid=f"b|r_tick|{g}|{s}")
            put(rx(v), ys + 0.9, s, f"b|r_ticklabel|{g}|{s}", fontsize=fs.PT["tiny"], color=fs.C["muted"], ha="center", va="top")
        for (i, j), v in B["r"][g].items():
            cv.plot([rx(v)], [ys], marker="o", ms=2.4, mew=0.5, mfc="white", mec=fs.C["muted"], ls="none", zorder=4,
                    gid=f"b|r|{g}|{i + 1}-{j + 1}")
        cv.plot([rx(B["mean_r"][g])] * 2, [ys - 1.1, ys + 0.7], color=fs.C["text"], lw=0.9, zorder=5, gid=f"b|r_mean|{g}")
    put(HX0, YKEY, "94 conditions in identifier order", "b|xlabel", fontsize=fs.PT["small"], ha="left", va="center")
    cax = fs.axes_mm(fig, 121.0, YKEY - 0.55, HX1 - 121.0, 1.1, H)
    cb = fig.colorbar(plt.cm.ScalarMappable(norm=norm, cmap=cmap), cax=cax, orientation="horizontal", extend="both",
                      extendfrac=0.06, ticks=[-30, 0, 30])
    cb.ax.set_xticklabels(["−30", "0", "+30"], fontsize=fs.PT["tiny"]); cb.ax.tick_params(length=1.5, width=0.4, pad=1)
    cb.outline.set_linewidth(0.3); cax.set_gid("b|colorbar")
    put(119.5, YKEY, "Condition pattern (percentage points)", "b|colorbar_label", fontsize=fs.PT["small"], ha="right", va="center")

    # ---------------------------------------------------------------- c and d: one row per array
    y_rows = [row_y(i) for i in range(len(T))]
    bottom = row_y(len(T) - 1) + PITCH / 2
    for i, (key, t) in enumerate(T.iterrows()):
        put(CX0 - 1.0, y_rows[i], t["array"], f"cd|array|{key}", fontsize=fs.PT["small"], ha="right", va="center")
    axc = fs.axes_mm(fig, CX0, ROW0, CW, bottom - ROW0, H); axc.set_gid("c|axes")
    axc.set_xlim(-10, 100); axc.set_ylim(bottom, ROW0)
    axc.axvline(0, color=fs.C["faint"], lw=0.4, zorder=0)
    for i, (key, t) in enumerate(T.iterrows()):
        y, hgt = y_rows[i], 2.3
        x = 0.0
        for comp, colr in (("molecular", fs.C["molecular"]), ("combination", fs.C["combination"]), ("noise", fs.C["noise"])):
            v = 100 * t[comp]
            axc.add_patch(Rectangle((x, y - hgt / 2), v, hgt, facecolor=colr, edgecolor="none", lw=0, gid=f"c|{comp}|{key}"))
            x += v
        axc.plot([100 * t.neither], [y], marker="o", ms=2.3, mew=0.5, mfc="white", mec=fs.C["muted"], ls="none",
                 gid=f"c|neither|{key}")
    axc.plot([100 * med["molecular"]] * 2, [ROW0 - 0.6, bottom], color=fs.C["molecular"], lw=0.5, ls=(0, (2.0, 1.5)),
             clip_on=False, zorder=3, gid="c|median_line|molecular")
    hy1, hy2 = YC + 4.4, YC + 6.9
    lab = lambda x, s, gid, y, size, colr, ha: axc.text(x, y, s, gid=gid, fontsize=size, color=colr, ha=ha, va="center",
                                                        clip_on=False)
    lab(-3.0, "neither shared", "c|name|neither", hy1, fs.PT["small"], fs.C["muted"], "right")
    lab(-3.0, neg(f"{100 * med['neither']:.0f}%"), "c|median|neither", hy2, fs.PT["base"], fs.C["muted"], "right")
    lab(5.0, "molecular", "c|name|molecular", hy1, fs.PT["small"], fs.C["molecular"], "left")
    lab(5.0, f"{100 * med['molecular']:.0f}% ({100 * ci[0]:.0f}–{100 * ci[1]:.0f}%)", "c|median|molecular", hy2,
        fs.PT["base"], fs.C["molecular"], "left")
    lab(66.0, "unattributed", "c|name|combination", hy1, fs.PT["small"], fs.C["muted"], "center")
    lab(66.0, f"{100 * med['combination']:.0f}%", "c|median|combination", hy2, fs.PT["base"], fs.C["text"], "center")
    lab(100.0, "noise", "c|name|noise", hy1, fs.PT["small"], fs.C["muted"], "right")
    lab(100.0, f"{100 * med['noise']:.1f}%", "c|median|noise", hy2, fs.PT["base"], fs.C["muted"], "right")
    axc.set_xticks([0, 25, 50, 75, 100]); axc.set_yticks([])
    axc.set_xlabel("Share of condition pattern (%)", labelpad=1.5)
    for sp in ("left", "right", "top"):
        axc.spines[sp].set_visible(False)
    axc.patch.set_visible(False)

    axd = fs.axes_mm(fig, DX0, ROW0, DW, bottom - ROW0, H); axd.set_gid("d|axes")
    axd.set_xlim(-10, 50); axd.set_ylim(bottom, ROW0)
    mm_per = DW / 60.0
    n_hal = int((T.group == "electrophilic halide").sum())
    hal_bottom = row_y(n_hal - 1) + PITCH / 2
    axd.add_patch(Rectangle((-10, ROW0), (W - 2.0 - DX0) / mm_per, hal_bottom - ROW0, facecolor=fs.C["band"],
                            edgecolor="none", zorder=0, clip_on=False, gid="d|band|electrophilic halide"))
    axd.axvline(0, color=fs.C["text"], lw=0.4, zorder=1)
    for i, (key, t) in enumerate(T.iterrows()):
        for which, dy, colr, tc in (("owner", -0.85, fs.C["owner"], fs.C["text"]), ("partner", 0.85, fs.C["partner"], fs.C["muted"])):
            v = 100 * t[f"{which}_share"]
            axd.add_patch(Rectangle((0, y_rows[i] + dy - 0.72), v, 1.44, facecolor=colr, edgecolor="none", lw=0, zorder=2,
                                    gid=f"d|{which}|{key}"))
            axd.text(max(v, 0) + 1.0, y_rows[i] + dy, t[which], fontsize=fs.PT["tiny"], color=tc, ha="left", va="center",
                     clip_on=False, gid=f"d|role|{which}|{key}")
    xr = (W - 2.0 - DX0) / mm_per - 10
    axd.text(xr, ROW0 + 0.6, f"electrophilic halide\nowns in {own['halide_owns']} of {own['halide_arrays']}",
             fontsize=fs.PT["tiny"], color=fs.C["muted"], ha="right", va="top", linespacing=1.15, clip_on=False,
             gid="d|group|electrophilic halide")
    axd.text(xr, hal_bottom + GROUP_GAP + 0.6, f"other {len(T) - n_hal} arrays", fontsize=fs.PT["tiny"], color=fs.C["muted"],
             ha="right", va="top", clip_on=False, gid="d|group|other")
    put(DX0 - 3.8, YC + 4.4, f"median owner fraction {own['median_owner_fraction']:.2f} · owner ≥ 2× partner in "
        f"{own['twice']} of {own['arrays']}", "d|summary|1", fontsize=fs.PT["tiny"], color=fs.C["muted"], ha="left", va="center")
    put(DX0 - 3.8, YC + 6.9, f"likelihood estimator names the same owner in {own['same_owner']} of {own['arrays']}",
        "d|summary|2", fontsize=fs.PT["tiny"], color=fs.C["muted"], ha="left", va="center")
    for j, (s, colr) in enumerate((("owner", fs.C["owner"]), ("partner", fs.C["partner"]))):
        x = W - 21.0 + 10.5 * j
        cv.add_patch(Rectangle((x, YC + 4.4 - 0.7), 1.4, 1.4, facecolor=colr, edgecolor="none", gid=f"d|legend|{s}"))
        put(x + 2.0, YC + 4.4, s, f"d|legend_label|{s}", fontsize=fs.PT["tiny"], ha="left", va="center")
    axd.set_xticks([0, 20, 40]); axd.set_yticks([])
    axd.set_xlabel("Share carried by each role (%)", labelpad=1.5)
    for sp in ("left", "right", "top"):
        axd.spines[sp].set_visible(False)
    axd.patch.set_visible(False)
    return fig


# ====================================================================== read back what was drawn
def readback(fig):
    fig.canvas.draw()
    rend = fig.canvas.get_renderer()
    px_mm = 25.4 / fig.dpi

    def bbox(a):
        e = a.get_window_extent(rend)
        return [round(e.x0 * px_mm, 3), round(H - e.y1 * px_mm, 3), round(e.x1 * px_mm, 3), round(H - e.y0 * px_mm, 3)]
    out = dict(width_mm=W, height_mm=H, images=[], rectangles=[], markers=[], lines=[], texts=[], polygons={})
    for ax in fig.axes:
        for im in ax.images:
            out["images"].append(dict(gid=im.get_gid(), data=np.asarray(im.get_array(), float).tolist(),
                                      clim=list(im.get_clim()), cmap=im.get_cmap().name, bbox_mm=bbox(im)))
        for p in ax.patches:
            g = p.get_gid()
            if not g:
                continue
            if isinstance(p, Polygon):
                out["polygons"].setdefault(g.rsplit("|", 1)[0], 0)
                out["polygons"][g.rsplit("|", 1)[0]] += 1
            elif isinstance(p, Rectangle):
                out["rectangles"].append(dict(gid=g, x=p.get_x(), y=p.get_y(), width=p.get_width(), height=p.get_height(),
                                              color=to_hex(p.get_facecolor()), bbox_mm=bbox(p)))
            else:
                out["rectangles"].append(dict(gid=g, kind=type(p).__name__, bbox_mm=bbox(p)))
        for ln in ax.lines:
            if ln.get_gid():
                rec = dict(gid=ln.get_gid(), x=np.asarray(ln.get_xdata(), float).tolist(),
                           y=np.asarray(ln.get_ydata(), float).tolist(), bbox_mm=bbox(ln))
                (out["markers"] if ln.get_marker() not in (None, "None", "") else out["lines"]).append(rec)
    for t in fig.findobj(Text):
        if t.get_visible() and t.get_text().strip():
            out["texts"].append(dict(gid=t.get_gid(), text=t.get_text(), size=t.get_fontsize(), bbox_mm=bbox(t)))
    return out


# ====================================================================== Source Data
def source_data(a_rows, B, T, med, own, ci):
    c_rows = T.reset_index()[["plot_row", "screen", "array", "group", "a1_rows", "conditions", "pairs", "product_forming",
                              "molecular", "combination", "noise", "neither"]].copy()
    for col in ("molecular", "combination", "noise", "neither"):
        c_rows[col] = 100 * c_rows[col]
    c_rows = c_rows.rename(columns={"molecular": "molecular_pct", "combination": "unattributed_pct",
                                    "noise": "noise_pct", "neither": "neither_shared_pct"})
    c_rows.insert(0, "row_type", "array")
    c_sum = pd.DataFrame([dict(row_type="median across the 15 arrays", molecular_pct=100 * med["molecular"],
                               unattributed_pct=100 * med["combination"], noise_pct=100 * med["noise"],
                               neither_shared_pct=100 * med["neither"]),
                          dict(row_type="95% interval of the molecular median, lower (approved output)", molecular_pct=100 * ci[0]),
                          dict(row_type="95% interval of the molecular median, upper (approved output)", molecular_pct=100 * ci[1])])
    d_rows = T.reset_index()[["plot_row", "screen", "array", "group", "role1", "role2", "share1", "share2", "owner",
                              "partner", "owner_share", "partner_share", "owner_fraction", "twice", "halide",
                              "halide_owns", "ml_owner_role", "ml_same_owner"]].copy()
    for col in ("share1", "share2", "owner_share", "partner_share"):
        d_rows[col] = 100 * d_rows[col]
    d_rows["ml_owner"] = [ROLE[k][int(r) - 1][0] for k, r in zip(d_rows.screen, d_rows.ml_owner_role)]
    d_rows = d_rows.drop(columns="ml_owner_role").rename(columns={
        "share1": "role1_share_pct", "share2": "role2_share_pct", "owner": "owner_role", "partner": "partner_role",
        "owner_share": "owner_share_pct", "partner_share": "partner_share_pct", "twice": "owner_at_least_2x_partner",
        "halide": "electrophilic_halide_role", "halide_owns": "halide_is_owner", "ml_owner": "likelihood_owner_role (approved output)",
        "ml_same_owner": "likelihood_same_owner"})
    d_rows.insert(0, "row_type", "array")
    d_sum = pd.DataFrame([dict(row_type="median owner fraction across the 15 arrays", owner_fraction=own["median_owner_fraction"]),
                          dict(row_type="arrays with owner >= 2x partner", count=own["twice"], of=own["arrays"]),
                          dict(row_type="arrays where the likelihood estimator names the same owner", count=own["same_owner"],
                               of=own["arrays"]),
                          dict(row_type="electrophilic-halide arrays where the halide is the owner", count=own["halide_owns"],
                               of=own["halide_arrays"])])
    readme = pd.DataFrame([
        dict(sheet="Fig1a", panel="a", rows="one per molecule drawn (4)", units="-", transformation=TRANSFORM["a"]),
        dict(sheet="Fig1b", panel="b", rows="one per heatmap cell (8 pairs x 94 conditions)",
             units="yields as fractions of 1; pattern in percentage points", transformation=TRANSFORM["b"]),
        dict(sheet="Fig1b_r", panel="b", rows="one per pair of rows within a group (2 x 6), then the two group means",
             units="Pearson r", transformation=TRANSFORM["b"]),
        dict(sheet="Fig1c", panel="c", rows="one per array (15), then the medians and the interval printed in the panel",
             units="% of the condition pattern", transformation=TRANSFORM["c"]),
        dict(sheet="Fig1d", panel="d", rows="one per array (15), then the four counts printed in the panel",
             units="% of the condition pattern; owner fraction as a fraction of 1", transformation=TRANSFORM["d"]),
        dict(sheet="(all)", panel="-", rows="-", units="-",
             transformation="Measurements: DATASET_analysis.csv.gz of the data release, rows of the 15 recurring-molecule "
                            "arrays with a numeric yield (177,308; Hiyama C46 excluded) and Hiyama C46 (74). Input hashes, "
                            "counts and the agreement with the analysis outputs: Figure1_SourceData_provenance.json."),
    ])
    ints = lambda df, cols: df.astype({c: "Int64" for c in cols if c in df})
    return dict(README=readme, Fig1a=a_rows, Fig1b=B["cells"], Fig1b_r=B["corr"],
                Fig1c=ints(pd.concat([c_rows, c_sum], ignore_index=True),
                           ("plot_row", "a1_rows", "conditions", "pairs", "product_forming")),
                Fig1d=ints(pd.concat([d_rows, d_sum], ignore_index=True), ("plot_row", "count", "of")))


def write_xlsx(sheets, path):
    from openpyxl.styles import Alignment, Font
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, col in enumerate(df.columns, start=1):
                vals = [str(col)] + [("%.6g" % v) if isinstance(v, float) else str(v) for v in df[col].tolist()]
                width = min(max(len(v) for v in vals) + 2, 60 if name != "README" else 110)
                ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = width
            if name == "README":
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    normalise_xlsx(path)


def normalise_xlsx(path):
    """No save time inside the file, so the same data give the same bytes and a delivered copy can be compared byte for byte."""
    import io
    import re
    import zipfile
    with zipfile.ZipFile(path) as src:
        items = [(i, src.read(i.filename)) for i in src.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"<dcterms:(?:created|modified)[^>]*>[^<]*</dcterms:(?:created|modified)>", b"", data)
                assert b"dcterms:created" not in data and b"dcterms:modified" not in data
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, info.external_attr
            z.writestr(zi, data)
    Path(path).write_bytes(buf.getvalue())


# ====================================================================== main
def main(out=None):
    paths = {k: (Path(out) / v.name if out else v) for k, v in OUT.items()}
    for p in paths.values():
        p.parent.mkdir(parents=True, exist_ok=True)
    sel = json.loads(APPROVED["Figure1b_example_selection.json"].read_text(encoding="utf-8"))
    frozen = pd.read_csv(APPROVED["Figure1b_example_residuals.csv"])
    J = json.loads(APPROVED["Anatomy.json"].read_text(encoding="utf-8"))
    approved_arrays = pd.read_csv(APPROVED["Anatomy_arrays.csv"])
    A1, A3 = public_units()
    a_rows, art, S = panel_a(A1, sel)
    B = panel_b(A1, sel, frozen, {S[L]: L for L in "BC"})
    T, med, own, ci, cd_info = panels_cd(A1, A3, approved_arrays, J)

    fig = draw(sel, a_rows, art, B, T, med, own, ci)
    rb = readback(fig)
    fig.savefig(paths["pdf"], metadata={"Creator": None, "Producer": None, "CreationDate": None})
    fig.savefig(paths["svg"], metadata={"Creator": None, "Date": None})
    fig.savefig(paths["png"], dpi=PNG_DPI, metadata={"Software": None})
    plt.close(fig)
    paths["readback"].write_text(json.dumps(rb, indent=1, ensure_ascii=False), encoding="utf-8")
    sheets = source_data(a_rows, B, T, med, own, ci)
    write_xlsx(sheets, paths["xlsx"])

    prov = dict(
        figure="Figure 1 | Condition knowledge is directional", generated_by=rel(__file__),
        inputs=dict(
            measurements=dict(file=rel(RAW), sha256=sha(RAW),
                              A1=dict(rows=N_A1, selection="screens of the 15 arrays (no process-development CMC screen), numeric yield, not Hiyama|1 C46"),
                              A3=dict(rows=N_A3, selection="Hiyama|1, condition C46 (duplicate of C43)")),
            approved_outputs={n: dict(file=rel(p), sha256=sha(p)) for n, p in APPROVED.items()},
            structures={L: dict(file=rel(CHEMDRAW / f"Fig1_a_{L}.svg"), sha256=art[L]["sha256"],
                                cdxml_sha256=sha(CHEMDRAW / f"Fig1_a_{L}.cdxml")) for L in "ABCD"},
            structure_spec=dict(file=rel(SPEC), sha256=sha(SPEC))),
        panels=dict(
            a=dict(records=len(a_rows), measurements_plotted=0, transformation=TRANSFORM["a"]),
            b=dict(B["info"], records=len(B["cells"]), correlations=len(B["corr"]) - 2, transformation=TRANSFORM["b"],
                   mean_r=B["mean_r"]),
            c=dict(cd_info, records=len(T), medians=med, molecular_median_interval=ci, transformation=TRANSFORM["c"]),
            d=dict(records=len(T), counts=own, transformation=TRANSFORM["d"])),
        taken_from_approved_outputs=["95% interval of the median molecular share (Anatomy.json share.ci; 2,000 hierarchical "
                                     "resamples of arrays and molecules)",
                                     "owner named by the likelihood estimator (Anatomy_arrays.csv ml_owner_role)"],
        outputs={k: dict(file=rel(OUT[k]), sha256=sha(p)) for k, p in paths.items() if k != "provenance"})
    paths["provenance"].write_text(json.dumps(prov, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"Figure 1: A1 {N_A1} rows, A3 {N_A3}; panel b max |diff| {B['info']['max_abs_diff_points']:.1e} points; "
          f"panels c-d max |diff| {max(cd_info['max_abs_diff'].values()):.1e}; wrote {', '.join(p.name for p in paths.values())}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out")
    main(ap.parse_args().out)
